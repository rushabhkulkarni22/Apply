from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

MAX_BYTES = 5 * 1024 * 1024
SKILLS = ('Python', 'SQL', 'PySpark', 'AWS', 'Azure', 'GCP', 'Spark', 'Airflow', 'Kafka',
          'Snowflake', 'Databricks', 'Docker', 'Kubernetes', 'Java', 'Scala', 'dbt', 'Terraform')


def parse_resume(filename, content):
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Upload a nonempty PDF or DOCX up to 5 MB')
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        if not content.startswith(b'%PDF-'):
            raise ValueError('This file is not a PDF')
        from pypdf import PdfReader
        doc = PdfReader(BytesIO(content))
        if doc.is_encrypted:
            raise ValueError('Remove the PDF password before uploading')
        if len(doc.pages) > 30:
            raise ValueError('Resume must have at most 30 pages')
        text = '\n'.join(page.extract_text() or '' for page in doc.pages)
    elif suffix == '.docx':
        try:
            with ZipFile(BytesIO(content)) as archive:
                files = archive.infolist()
                if len(files) > 1000 or sum(f.file_size for f in files) > 20 * 1024 * 1024:
                    raise ValueError('Word document is too large when expanded')
                if 'word/document.xml' not in archive.namelist():
                    raise ValueError('This file is not a DOCX document')
                if any('vbaproject' in f.filename.lower() for f in files):
                    raise ValueError('Documents containing macros are not accepted')
                for f in files:
                    if f.filename.endswith(('.xml', '.rels')):
                        xml = archive.read(f)
                        if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml or b'TargetMode="External"' in xml and b'oleObject' in xml:
                            raise ValueError('Unsupported active document content')
        except BadZipFile:
            raise ValueError('This file is not a DOCX document')
        from docx import Document
        document = Document(BytesIO(content))
        text = '\n'.join([p.text for p in document.paragraphs] +
                         [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    else:
        raise ValueError('Only PDF and DOCX are supported; convert legacy DOC to DOCX')
    if len(text.strip()) < 30:
        raise ValueError('No readable resume text found. Upload a text PDF or DOCX; scanned PDFs need OCR first')
    if len(text) > 150000:
        raise ValueError('Resume text exceeds the supported limit')
    skills = [s for s in SKILLS if re.search(r'(?<!\w)' + re.escape(s) + r'(?!\w)', text, re.I)]
    if re.search(r'amazon web services', text, re.I) and 'AWS' not in skills:
        skills.append('AWS')
    email = re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', text)
    return {'text': text, 'skills': skills, 'email': email[0] if email else '',
            'suggested_name': text.strip().splitlines()[0][:150], 'parser_version': '1',
            'message': 'Review these suggestions. Work experience and personal answers require your confirmation.'}
